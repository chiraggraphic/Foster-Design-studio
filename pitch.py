"""Client pitch / design presentation builder. Produces an editable .pptx (python-pptx) with speaker notes."""
import asyncio, datetime, io
import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
import mockups, tools

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
THEMES = {"light": dict(bg="FFFFFF", fg="1A1A1A", mute="5F6368", card="F4F4F5"), "dark": dict(bg="141414", fg="F5F5F5", mute="A8A29E", card="232323")}
HEAD, BODY = "Cambria", "Calibri"
DEFAULT_STEPS = ["Review the concepts and share your feedback", "Choose the direction to develop", "Refinement rounds as agreed in the quotation", "Final files and handover"]
C = RGBColor.from_string
lines = lambda t: [l.strip() for l in (t or "").splitlines() if l.strip()]

def fit(ls, w, h, start, lo=14, bold=False):
    """Pick the largest font size that fits the box; at the minimum size, trim text with an ellipsis instead of overflowing."""
    cw = .55 if bold else .5
    def need(s):
        cpl = max(1, int(w * 72 / (s * cw))); n = sum(max(1, -(-len(l) // cpl)) for l in ls)
        return (n * s * 1.2 + (len(ls) - 1) * s * .5) / 72
    for s in range(start, lo - 1, -1):
        if need(s) <= h: return s, ls
    cpl = max(1, int(w * 72 / (lo * cw))); room = max(1, int(h * 72 / (lo * 1.2))); out = []
    for l in ls:
        k = max(1, -(-len(l) // cpl))
        if k > room:
            if room * cpl > 8: out.append(l[:room * cpl - 1].rstrip() + "…")
            break
        out.append(l); room -= k
    return lo, out or [ls[0][:cpl - 1] + "…"]

class Deck:
    def __init__(s, theme, accent, agency, project):
        s.p = Presentation(); s.p.slide_width, s.p.slide_height = Inches(13.333), Inches(7.5)
        s.t, s.ac, s.agency, s.project, s.n = THEMES[theme], accent, agency, project, 0

    def slide(s, notes="", foot=True):
        sl = s.p.slides.add_slide(s.p.slide_layouts[6]); f = sl.background.fill; f.solid(); f.fore_color.rgb = C(s.t["bg"]); s.n += 1
        if notes: sl.notes_slide.notes_text_frame.text = notes
        if foot:
            s.text(sl, .6, 7.0, 9, .3, f"{s.agency}  |  {s.project}", 10, s.t["mute"])
            s.text(sl, 11.7, 7.0, 1.0, .3, str(s.n), 10, s.t["mute"], align=PP_ALIGN.RIGHT)
        return sl

    def shape(s, sl, kind, x, y, w, h, fill, round_=None):
        sp = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); sp.fill.solid(); sp.fill.fore_color.rgb = C(fill)
        sp.line.fill.background(); sp.shadow.inherit = False
        if round_: sp.adjustments[0] = round_
        return sp

    def card(s, sl, x, y, w, h, fill=None): return s.shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill or s.t["card"], .04)

    def text(s, sl, x, y, w, h, ls, size, color=None, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lo=None):
        ls = [ls] if isinstance(ls, str) else list(ls)
        if lo: size, ls = fit(ls, w - .2, h, size, lo, bold)
        tf = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, l in enumerate(ls):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(size * .5) if i < len(ls) - 1 else Pt(0)
            r = p.add_run(); r.text = l; r.font.size = Pt(size); r.font.bold = bold; r.font.name = font; r.font.color.rgb = C(color or s.t["fg"])

    def label(s, sp, txt, size, color="FFFFFF"):
        tf = sp.text_frame; tf.margin_left = tf.margin_right = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = C(color)

    def pic(s, sl, im, x, y, w, h, pad=0):
        im = im.copy(); im.thumbnail((2000, 2000)); keep = im.mode == "RGBA" and im.getchannel("A").getextrema()[0] < 255; buf = io.BytesIO()
        (im if keep else im.convert("RGB")).save(buf, "PNG" if keep else "JPEG", **({} if keep else {"quality": 88}))
        bw, bh = w - 2 * pad, h - 2 * pad; k = min(bw / im.width, bh / im.height); pw, ph = im.width * k, im.height * k
        sl.shapes.add_picture(buf, Inches(x + pad + (bw - pw) / 2), Inches(y + pad + (bh - ph) / 2), Inches(pw), Inches(ph))

    def head(s, sl, title, kicker=""):
        if kicker: s.text(sl, .6, .4, 8, .35, kicker.upper(), 14, s.ac, True)
        s.text(sl, .6, .72, 12.1, .85, title, 36, None, True, HEAD, lo=28)

    def row(s, sl, x, y, w, h, n, txt, size=18):
        s.card(sl, x, y, w, h); o = s.shape(sl, MSO_SHAPE.OVAL, x + .2, y + (h - .55) / 2, .55, .55, s.ac); s.label(o, str(n), 16)
        s.text(sl, x + .95, y, w - 1.1, h, txt, size, anchor=MSO_ANCHOR.MIDDLE, lo=14)

    # ---- slides ----
    def title(s, client, project, date, im, notes):
        sl = s.slide(notes, foot=False)
        s.text(sl, .8, 1.0, 6.5, .4, "DESIGN PROPOSAL", 14, s.ac, True)
        s.text(sl, .8, 1.6, 7.0, 2.6, project, 54, None, True, HEAD, lo=36)
        s.text(sl, .8, 4.4, 7.0, .6, f"Prepared for {client}", 24, s.t["mute"], lo=16)
        s.text(sl, .8, 6.2, 7.0, .8, [s.agency, date], 14, s.t["mute"])
        s.card(sl, 8.4, .7, 4.4, 6.1); s.pic(sl, im, 8.4, .7, 4.4, 6.1, pad=.3)

    def about(s, about, services):
        sl = s.slide("Introduce the studio in one minute: who you are, what you do, and why the client can trust you with this project.")
        s.head(sl, f"About {s.agency}", "Who we are")
        s.card(sl, .6, 1.8, 5.6, 4.9); s.text(sl, .9, 2.0, 5.0, 4.5, about or [s.agency], 18, None, False, BODY, lo=14)
        for i, v in enumerate(services[:6]): s.row(sl, 6.6, 1.8 + i * .83, 6.1, .73, i + 1, v, 18)

    def brief(s, brief, objectives):
        sl = s.slide("Show you listened. Read the brief back in your own words and confirm the objectives before showing any design.")
        s.head(sl, "Understanding the brief", "Where we started")
        if brief and objectives:
            s.card(sl, .6, 1.8, 5.8, 4.9); s.text(sl, .9, 2.0, 5.2, .5, "The brief", 20, s.ac, True); s.text(sl, .9, 2.6, 5.2, 3.9, lines(brief) or [brief], 18, lo=14)
            for i, v in enumerate(objectives[:4]): s.row(sl, 6.8, 1.8 + i * 1.25, 5.9, 1.1, i + 1, v, 18)
        elif brief: s.card(sl, .6, 1.8, 12.1, 4.9); s.text(sl, 1.0, 2.1, 11.3, 4.3, lines(brief), 20, lo=14)
        else:
            for i, v in enumerate(objectives[:4]): s.row(sl, .6, 1.8 + i * 1.25, 12.1, 1.1, i + 1, v, 20)

    def idea(s, text, im):
        sl = s.slide("Explain the thinking: why this direction suits the client's audience and goals. Keep it to the story, not the software.")
        s.head(sl, "The idea", "Design rationale")
        s.text(sl, .6, 1.6, 1.5, 1.5, "“", 120, s.ac, True, HEAD); s.text(sl, 1.4, 2.4, 5.9, 4.3, lines(text) or [text], 22, lo=16)
        s.card(sl, 7.8, 1.8, 4.9, 4.9); s.pic(sl, im, 7.8, 1.8, 4.9, 4.9, pad=.3)

    def palette(s, cols):
        sl = s.slide("Colours carry the brand. Point out the main colour, the supporting ones, and which text colour stays readable on each.")
        s.head(sl, "Colour palette", "Brand colours"); n = len(cols); g = .25; w = (12.1 - g * (n - 1)) / n
        for i, c in enumerate(cols):
            x = .6 + i * (w + g); sp = s.shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.8, w, 3.3, c["hex"][1:].upper(), .06)
            s.text(sl, x + .1, 1.9, w - .2, .9, "Aa", 40, c["text_on"][1:].upper(), True, HEAD)
            s.text(sl, x, 5.3, w, .4, c["hex"].upper(), 18, None, True); s.text(sl, x, 5.75, w, .6, [f"{c['percent']:g}% of the design", f"Best text: {'white' if c['text_on'] == '#ffffff' else 'black'}"], 12, s.t["mute"])

    def showcase(s, i, n, im, caption):
        sl = s.slide("Present this design: what problem it solves, then the main decisions (colour, type, layout). Pause for reactions before moving on.")
        s.head(sl, f"Design {i}" if n > 1 else "The design", f"{i} of {n}" if n > 1 else "Concept")
        s.card(sl, .6, 1.7, 12.1, 4.7); s.pic(sl, im, .6, 1.7, 12.1, 4.7, pad=.25)
        if caption: s.text(sl, .6, 6.5, 12.1, .45, caption, 16, s.t["mute"], lo=14)

    def inuse(s, imgs):
        """Two big mockups per slide (or one large + two small when there are three)."""
        groups = [imgs[i:i + 2] for i in range(0, len(imgs), 2)] if len(imgs) != 3 else [imgs]
        for gi, grp in enumerate(groups):
            sl = s.slide("Show the design in real situations so the client can picture it. Mockups here are illustrative previews.")
            s.head(sl, "In use", "Mockups" + (f" {gi + 1} of {len(groups)}" if len(groups) > 1 else ""))
            cells = ([(.6, 1.7, 7.4, 5.15), (8.2, 1.7, 4.5, 2.5), (8.2, 4.35, 4.5, 2.5)] if len(grp) == 3
                     else [(.6, 1.7, 5.95, 5.0), (6.75, 1.7, 5.95, 5.0)] if len(grp) == 2 else [(.6, 1.7, 12.1, 5.0)])
            for im, (x, y, w, h) in zip(grp, cells): s.card(sl, x, y, w, h); s.pic(sl, im, x, y, w, h, pad=.1)

    def invest(s, items, gst, terms):
        sl = s.slide("Walk through each line, then confirm payment terms and what is included in the revision rounds.")
        s.head(sl, "Investment", "Quotation"); cur = tools.CUR; sub = sum(q * r for _, q, r in items); g = sub * gst / 100
        rows = [("Description", "Qty", "Rate", "Amount")] + [(d, f"{q:g}", f"{cur} {r:,.2f}", f"{cur} {q * r:,.2f}") for d, q, r in items] + [("", "", "Subtotal", f"{cur} {sub:,.2f}")]
        if g: rows.append(("", "", f"GST ({gst:g}%)", f"{cur} {g:,.2f}"))
        rows.append(("", "", "Total", f"{cur} {sub + g:,.2f}")); n = len(items)
        tb = sl.shapes.add_table(len(rows), 4, Inches(.6), Inches(1.75), Inches(12.1), Inches(.42 * len(rows))).table
        for i, w in enumerate((6.1, 1.2, 2.4, 2.4)): tb.columns[i].width = Inches(w)
        for ri, row in enumerate(rows):
            tb.rows[ri].height = Inches(.42)
            for ci, v in enumerate(row):
                c = tb.cell(ri, ci); c.fill.solid(); c.fill.fore_color.rgb = C(s.ac if ri == 0 else (s.t["card"] if ri > n or ri % 2 == 0 else s.t["bg"])); c.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if ci else PP_ALIGN.LEFT; r = p.add_run(); r.text = v
                r.font.size = Pt(14); r.font.name = BODY; r.font.bold = ri == 0 or ri > n; r.font.color.rgb = C("FFFFFF" if ri == 0 else s.t["fg"])
        y = 1.75 + .42 * len(rows) + .3
        if terms and y < 6.4: s.text(sl, .6, y, 12.1, 6.9 - y, lines(terms), 14, s.t["mute"], lo=14)

    def timeline(s, stages):
        sl = s.slide("Set expectations on pace. Each stage ends with a client checkpoint, so feedback timing affects the total.")
        s.head(sl, "Timeline", "How it will run"); n = len(stages); step = 12.1 / n; day = 0
        s.shape(sl, MSO_SHAPE.RECTANGLE, .6 + step / 2, 2.52, 12.1 - step, .06, s.t["card"])
        for i, (name, d) in enumerate(stages):
            cx = .6 + step * (i + .5); o = s.shape(sl, MSO_SHAPE.OVAL, cx - .35, 2.25, .7, .7, s.ac); s.label(o, str(i + 1), 20)
            s.text(sl, cx - step / 2 + .05, 3.25, step - .1, 1.1, name, 20, None, True, align=PP_ALIGN.CENTER, lo=14)
            s.text(sl, cx - step / 2 + .05, 4.45, step - .1, .4, f"{d:g} day{'s' if d != 1 else ''}", 16, s.t["mute"], align=PP_ALIGN.CENTER)
            s.text(sl, cx - step / 2 + .05, 4.85, step - .1, .4, f"Day {day + 1 if d else day:g}–{day + d:g}", 14, s.ac, True, align=PP_ALIGN.CENTER); day += d
        s.card(sl, .6, 5.6, 12.1, 1.1); s.text(sl, .9, 5.6, 11.5, 1.1, f"Total: about {day:g} working days, counted from the day we receive your approval and materials", 20, None, True, anchor=MSO_ANCHOR.MIDDLE, lo=14)

    def close(s, steps, contact, client):
        sl = s.slide("Close with a clear ask: agree the direction and the next date. Leave the contact slide up during questions.")
        s.head(sl, "Next steps", "Let's move forward")
        for i, v in enumerate(steps[:4]): s.row(sl, .6, 1.8 + i * 1.25, 7.2, 1.1, i + 1, v, 18)
        s.card(sl, 8.2, 1.8, 4.5, 4.85, "1A1A1A"); s.text(sl, 8.5, 2.1, 3.9, 1.5, "Let's get started", 32, s.ac, True, HEAD, lo=24)
        s.text(sl, 8.5, 3.7, 3.9, 2.7, [f"Thank you, {client}."] + lines(contact), 16, "F5F5F5", lo=14)

def llm(prompt):
    try: return httpx.post(f"{tools.LLM_URL}/api/generate", json={"model": tools.LLM_MODEL, "prompt": prompt, "stream": False}, timeout=180).json()["response"].strip()
    except Exception: return ""

def parse_items(txt):
    out = []
    for l in lines(txt):
        a = [x.strip() for x in l.split("|")]
        try: out.append((a[0], float(a[1]) if len(a) > 1 and a[1] else 1.0, float(a[2]) if len(a) > 2 else 0.0))
        except ValueError: raise HTTPException(400, f"Could not read this line: '{l}'. Use: description | quantity | rate")
    return out

def make_deck(p, designs, logo):
    full = p["mode"] == "full"; src = logo if (logo is not None and p["dtype"] == "logo") else designs[0]
    d = Deck(p["theme"], p["accent"], tools.AGENCY, p["project"]); today = datetime.date.today().strftime("%d %B %Y")
    objectives, rationale = lines(p["objectives"]), p["rationale"].strip()
    if p["ai_write"] and full and p["brief"].strip():
        ctx = f"Client: {p['client']}. Project: {p['project']}. Brief and notes: {p['brief'][:1500]}"
        if not objectives:
            objectives = [l.lstrip("-*•0123456789.) ").strip() for l in lines(llm(f"{ctx}\nList 4 short, specific project objectives, one per line. No numbering, no introduction."))][:4]
        if not rationale:
            rationale = llm(f"{ctx}\nWrite a 3 to 4 sentence design rationale for a client presentation. Do not invent facts about the client. Plain text only.")
    d.title(p["client"], p["project"], today, logo or designs[0], "Welcome the client, state the purpose of the meeting and how long you need.")
    if full:
        about, services = p["about"].strip() or tools.E("AGENCY_ABOUT", ""), lines(p["services"]) or [x.strip() for x in tools.E("AGENCY_SERVICES", "").split(",") if x.strip()]
        if about or services: d.about(lines(about), services)
        if p["brief"].strip() or objectives: d.brief(p["brief"], objectives)
        if rationale: d.idea(rationale, designs[0])
        if p["palette"]: d.palette(tools.palette_of(logo or designs[0], 5))
    caps = lines(p["captions"])
    for i, im in enumerate(designs): d.showcase(i + 1, len(designs), im, caps[i] if i < len(caps) else "")
    if p["mockups"]:
        mode = "logo" if p["dtype"] == "logo" else "art"; imgs = [mockups.render(k, src, None, "", "", mode) for k in mockups.SETS[p["dtype"]]]
        d.inuse(imgs)
    items = parse_items(p["items"])
    if full and items: d.invest(items, p["gst_percent"], p["terms"])
    stages = []
    for l in lines(p["timeline"]):
        a = [x.strip() for x in l.split("|")]
        try: stages.append((a[0], float(a[1]) if len(a) > 1 and a[1] else 0.0))
        except ValueError: raise HTTPException(400, f"Could not read timeline line: '{l}'. Use: stage | days")
    if full and stages: d.timeline(stages[:6])
    d.close(lines(p["next_steps"]) or DEFAULT_STEPS, tools.CONTACT, p["client"])
    d.p.core_properties.title = f"{p['project']} - {p['client']}"; d.p.core_properties.author = tools.AGENCY
    buf = io.BytesIO(); d.p.save(buf); return buf.getvalue(), d.n

def make_router(auth):
    r = APIRouter(prefix="/api/pitch", dependencies=[Depends(auth)])

    @r.post("/build")
    async def build(files: list[UploadFile] = File(...), logo: UploadFile | None = File(None), client: str = Form(...), project: str = Form(...),
                    mode: str = Form("full"), dtype: str = Form("logo"), theme: str = Form("light"), accent: str = Form("#f26b21"),
                    brief: str = Form(""), objectives: str = Form(""), rationale: str = Form(""), ai_write: str = Form("0"), mockups: str = Form("1"),
                    palette: str = Form("1"), captions: str = Form(""), about: str = Form(""), services: str = Form(""), items: str = Form(""),
                    gst_percent: float = Form(0), timeline: str = Form(""), terms: str = Form(""), next_steps: str = Form("")):
        if not 1 <= len(files) <= 20: raise HTTPException(400, "Upload 1-20 design images")
        if mode not in ("full", "quick") or theme not in THEMES or dtype not in mockups_sets(): raise HTTPException(400, "Bad mode, theme or design type")
        if len(parse_items(items)) > 8: raise HTTPException(400, "Put at most 8 line items on the pitch slide (use the Quotation PDF for longer lists)")
        acc = "%02X%02X%02X" % tools.color(accent)
        ims = [tools.load(await f.read()) for f in files]; lg = tools.load(await logo.read()) if logo is not None and logo.filename else None
        cut = lambda t, n=2500: (t or "")[:n]
        p = dict(client=cut(client, 120), project=cut(project, 120), mode=mode, dtype=dtype, theme=theme, accent=acc, brief=cut(brief), objectives=cut(objectives),
                 rationale=cut(rationale), ai_write=ai_write == "1", mockups=mockups == "1", palette=palette == "1", captions=cut(captions), about=cut(about),
                 services=cut(services), items=cut(items), gst_percent=max(0, min(gst_percent, 100)), timeline=cut(timeline), terms=cut(terms), next_steps=cut(next_steps))
        data, n = await asyncio.to_thread(make_deck, p, ims, lg)
        return Response(data, media_type=PPTX, headers={"X-Slides": str(n)})
    return r

def mockups_sets(): return mockups.SETS
